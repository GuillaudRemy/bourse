"""Parcours complet par l'API, sans internet : valeur importee, cours CSV, analyse, fondamentaux."""
from fastapi.testclient import TestClient

from backend.main import app


def test_parcours_complet(csv_cours):
    with TestClient(app) as client:
        assert client.get("/health").json()["status"] == "ok"

        valeur = client.post("/valeurs", json={"symbole": "test", "source": "import", "nom": "Test"}).json()
        vid = valeur["id"]
        assert valeur["symbole"] == "TEST"
        assert client.post("/valeurs", json={"symbole": "TEST", "source": "import"}).status_code == 409

        r = client.post(f"/valeurs/{vid}/import/cours", files={"fichier": ("cours.csv", csv_cours)})
        assert r.status_code == 200, r.text
        assert r.json()["cours"] > 6000

        r = client.post(f"/valeurs/{vid}/import/fondamentaux", files={
            "fichier": ("f.csv", "Année;BPA;Bénéfice net\n2015;2,5;1 200M\n2020;3,9;1,9Md\n".encode()),
        })
        assert r.status_code == 200, r.text
        assert client.put(f"/valeurs/{vid}/fondamentaux/2024", json={"bpa": 5.1}).json()["source"] == "manuel"

        analyse = client.get(f"/valeurs/{vid}/analyse").json()
        assert abs(analyse["regression"]["pente"] - 0.10) < 0.02
        assert analyse["base_ajustee"] is False          # pas de colonne ajustee dans le CSV
        assert len(analyse["fiabilite"]["verifications"]) == 3
        technique = client.get(f"/valeurs/{vid}/technique", params={"unite": "semaine", "mm": "10,30"}).json()
        assert len(technique["jours"]) == len(technique["rsi"]) == len(technique["moyennes"]["30"])
        assert len(technique["jours"]) < len(analyse["series"]["jours"]) / 4          # bougies hebdomadaires
        assert all(v is None or 0 <= v <= 100 for v in technique["stochastique_k"])

        ligne = next(l for l in client.get("/valeurs/synthese").json() if l["valeur"]["id"] == vid)
        assert ligne["zone"]

        mauvais = client.post(f"/valeurs/{vid}/import/cours", files={"fichier": ("x.csv", b"a,b\n1,2")})
        assert mauvais.status_code == 400

        assert client.delete(f"/valeurs/{vid}").status_code == 204
        assert client.get(f"/valeurs/{vid}").status_code == 404


def test_collage_tableur_annees_en_colonnes():
    """Forme habituelle d'un tableur : une ligne par poste, les annees en colonnes, tabulations."""
    texte = ("Poste\t2018\t2019\t2020\n"
             "Chiffre d'affaires\t1 200,5\t1 310\t1 290\n"
             "Bénéfice net (part du groupe)\t100\t110\t95\n"
             "BPA dilué\t2,10\t2,30\t1,99\n"
             "Ligne inconnue\t1\t2\t3\n")
    with TestClient(app) as client:
        vid = client.post("/valeurs", json={"symbole": "COLLE", "source": "import"}).json()["id"]
        r = client.post(f"/valeurs/{vid}/import/fondamentaux/texte", json={"texte": texte})
        assert r.status_code == 200, r.text
        lignes = {f["annee"]: f for f in client.get(f"/valeurs/{vid}/fondamentaux").json()}
        assert lignes[2018]["chiffre_affaires"] == 1200.5
        assert lignes[2019]["benefice_net"] == 110
        assert lignes[2020]["bpa"] == 1.99
        assert lignes[2020]["sources"]["bpa"] == "import"


def test_pdf_propose_puis_valide():
    from pathlib import Path
    pdf = (Path(__file__).parent / "donnees" / "rapport_exemple.pdf").read_bytes()
    with TestClient(app) as client:
        vid = client.post("/valeurs", json={"symbole": "PDF", "source": "import"}).json()["id"]
        proposition = client.post(f"/valeurs/{vid}/import/pdf", files={"fichier": ("r.pdf", pdf)}).json()
        assert proposition["multiplicateur"] == 1e6
        postes = {p["cle"]: p["valeurs"] for p in proposition["postes"]}
        assert postes["chiffre_affaires"] == {"2024": 27058.0, "2023": 27607.6}
        assert postes["bpa"] == {"2024": 5.72, "2023": 5.33}
        assert client.get(f"/valeurs/{vid}/fondamentaux").json() == []      # rien d'enregistre

        lot = {"annees": {"2024": {"chiffre_affaires": 27058e6, "bpa": 5.72, "charges_interets": -412.3e6}}}
        assert client.post(f"/valeurs/{vid}/import/lot", json=lot).status_code == 200
        ligne = client.get(f"/valeurs/{vid}/fondamentaux").json()[0]
        assert ligne["chiffre_affaires"] == 27058e6 and ligne["sources"]["bpa"] == "pdf"
        assert ligne["charges_interets"] == 412.3e6             # depense stockee en positif


def test_actualisation_en_tache_de_fond():
    import time
    with TestClient(app) as client:
        client.post("/valeurs", json={"symbole": "FOND1", "source": "import"})
        etat = client.post("/valeurs/actualiser").json()
        assert etat["total"] >= 1
        for _ in range(600):              # d'autres tests ont pu creer des valeurs Yahoo : appels reseau
            etat = client.get("/valeurs/actualisation").json()
            if not etat["en_cours"]:
                break
            time.sleep(0.1)
        assert not etat["en_cours"] and etat["faits"] == etat["total"]
        assert etat["fin"]


def test_indice_en_points(csv_cours):
    """Un indice se lit en points ; un ETF ou une action, dans sa devise."""
    with TestClient(app) as client:
        def creer(symbole, type_):
            return client.post("/valeurs", json={"symbole": symbole, "source": "import", "type": type_, "devise": "EUR"}).json()
        indice, etf = creer("IDXTEST", "indice"), creer("ETFTEST", "etf")
        for v in (indice, etf):
            client.post(f"/valeurs/{v['id']}/import/cours", files={"fichier": ("c.csv", csv_cours)})
        assert client.get(f"/valeurs/{indice['id']}/analyse").json()["devise_affichee"] == "pts"
        assert client.get(f"/valeurs/{etf['id']}/analyse").json()["devise_affichee"] == "EUR"


def test_exports_excel(csv_cours):
    import io

    import openpyxl
    with TestClient(app) as client:
        vid = client.post("/valeurs", json={"symbole": "EXPORT", "source": "import", "nom": "Export SA"}).json()["id"]
        client.post(f"/valeurs/{vid}/import/cours", files={"fichier": ("c.csv", csv_cours)})
        client.put(f"/valeurs/{vid}/fondamentaux/2024", json={"bpa": 5.0, "chiffre_affaires": 1e9})
        r = client.get(f"/valeurs/{vid}/export.xlsx")
        assert r.status_code == 200 and "EXPORT" in r.headers["content-disposition"]
        classeur = openpyxl.load_workbook(io.BytesIO(r.content))
        assert {"Résumé", "Cours mensuels", "Comptes annuels"} <= set(classeur.sheetnames)
        assert classeur["Cours mensuels"].max_row > 200
        assert client.get("/portefeuille/export.xlsx").status_code == 200


def test_a_propos_et_diagnostic_sans_donnees_personnelles():
    import io
    import json
    import zipfile
    with TestClient(app) as client:
        client.put("/reglages/contact_sec", json={"valeur": "personne@exemple.fr"})
        infos = client.get("/a-propos").json()
        assert infos["version"] and "valeur" in infos["enregistrements"]
        r = client.get("/diagnostic.zip")
        archive = zipfile.ZipFile(io.BytesIO(r.content))
        contenu = json.loads(archive.read("diagnostic.json"))
        assert contenu["reglages"]["contact_sec"] == "(renseigné)"
        assert "personne@exemple.fr" not in archive.read("diagnostic.json").decode()


def test_sauvegarde_externe_et_restauration(tmp_path):
    externe = tmp_path / "Drive"
    externe.mkdir()
    with TestClient(app) as client:
        client.put("/reglages/dossier_sauvegarde_externe", json={"valeur": str(externe)})
        client.post("/valeurs", json={"symbole": "AVANT", "source": "import", "nom": "Avant"})
        assert "dossier externe" in client.post("/sauvegardes").json()["message"]
        copies = client.get("/sauvegardes").json()["copies"]
        externe_copie = next(c for c in copies if c["emplacement"] == "externe")
        assert list((externe / "Bourse").glob("bourse-*.db")) and not list((externe / "Bourse").glob("*.partiel"))

        client.post("/valeurs", json={"symbole": "APRES", "source": "import", "nom": "Après"})
        assert client.post("/sauvegardes/restaurer", json={"chemin": "/etc/passwd"}).status_code == 404
        r = client.post("/sauvegardes/restaurer", json={"chemin": externe_copie["chemin"]})
        assert r.status_code == 200, r.text
        symboles = {v["symbole"] for v in client.get("/valeurs").json()}
        assert "AVANT" in symboles and "APRES" not in symboles
        assert any(c["nom"].startswith("avant-restauration-") for c in client.get("/sauvegardes").json()["copies"])


def test_suppression_groupee():
    with TestClient(app) as client:
        ids = [client.post("/valeurs", json={"symbole": f"SUPPR{i}", "source": "import", "nom": f"Suppr {i}"}).json()["id"]
               for i in range(3)]
        compte = client.post("/portefeuille/comptes", json={"nom": "Compte suppr", "type": "CTO"}).json()
        client.post("/portefeuille/transactions", json={"compte_id": compte["id"], "valeur_id": ids[0], "jour": "2024-03-01",
                                                        "type": "achat", "quantite": 1, "prix": 10, "frais": 0})
        apercu = client.post("/valeurs/supprimer", json={"ids": ids[:2]}).json()
        assert apercu == {"noms": ["Suppr 0", "Suppr 1"], "operations": 1}
        assert {v["symbole"] for v in client.get("/valeurs").json()} >= {"SUPPR0", "SUPPR1"}   # rien supprime
        assert client.post("/valeurs/supprimer", json={"ids": ids[:2], "confirmer": True}).json()["supprimees"] == 2
        restants = {v["symbole"] for v in client.get("/valeurs").json()}
        assert "SUPPR0" not in restants and "SUPPR1" not in restants and "SUPPR2" in restants
        operations = client.get("/portefeuille/transactions").json()
        assert any(t["compte_id"] == compte["id"] and t["valeur_id"] is None for t in operations)   # operation gardee


def test_reponses_brutes_purgees(tmp_path, monkeypatch):
    """Reponses gardees par jour ; au-dela de 30 jours, ou du plafond de taille, les plus anciennes partent."""
    from datetime import date, timedelta

    from backend.core import brut
    from backend.core import config as module_config
    monkeypatch.setattr(module_config.config, "DB_NAME", str(tmp_path / "b.db"))
    monkeypatch.setattr(module_config.config, "CONSERVER_BRUT", True)
    brut.conserver("data.sec.gov", "/api/xbrl/companyfacts/CIK0000320193.json", b'{"a": 1}' * 100)
    vieux = brut.dossier() / (date.today() - timedelta(days=31)).isoformat()
    vieux.mkdir()
    (vieux / "x.gz").write_bytes(b"x")
    assert brut.resume()["fichiers"] == 2
    brut.purger()
    assert not vieux.exists() and brut.resume()["fichiers"] == 1
    monkeypatch.setattr(brut, "PLAFOND", 0)
    brut.purger()                                  # le jour en cours est toujours garde
    assert brut.resume()["fichiers"] == 1


def test_reglages_en_menus():
    with TestClient(app) as client:
        reglages = {r["cle"]: r for r in client.get("/reglages").json()}
        assert ["20", "20"] in reglages["valeurs_par_page"]["choix"]
        # Valeur par defaut toujours proposee, sinon le menu afficherait autre chose que le reglage reel.
        for r in reglages.values():
            if r["choix"]:
                assert r["defaut"] in [v for v, _ in r["choix"]], r["cle"]
        assert client.put("/reglages/canal_beta", json={"valeur": "peut-etre"}).status_code == 422
        assert client.put("/reglages/valeurs_par_page", json={"valeur": "33"}).status_code == 422
        assert client.put("/reglages/frais_fixe", json={"valeur": "1,23"}).status_code == 200    # « Autre… »
        assert client.put("/reglages/tranche_marginale", json={"valeur": ""}).status_code == 200
        client.put("/reglages/frais_fixe", json={"valeur": ""})
